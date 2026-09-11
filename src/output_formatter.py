"""
output_formatter.py — Structured output formatter for HAZOP NLP system.

Produces:
  1. JSON  — machine-readable full dataset
  2. HTML  — print-ready HAZOP worksheet with colour-coded risk ratings
  3. Excel — .xlsx with multiple sheets (Summary, HAZOP Table, NER Entities)
  4. CSV   — flat export for further analysis

All functions accept list[dict] (rows_to_dicts() output from hazop_engine).
"""

from __future__ import annotations

import json
import csv
import io
from pathlib import Path
from datetime import datetime
from typing import Optional

from loguru import logger
from config import OUTPUT_DIR

# ── Optional heavy imports ────────────────────────────────────────────────────
try:
    import pandas as pd
    _PANDAS = True
except ImportError:
    _PANDAS = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter
    _OPENPYXL = True
except ImportError:
    _OPENPYXL = False


# ══════════════════════════════════════════════════════════════════════════════
# Risk Colour Scheme
# ══════════════════════════════════════════════════════════════════════════════

RISK_COLOURS = {
    "Critical": {"bg": "#C0392B", "fg": "#FFFFFF", "hex": "C0392B"},
    "High":     {"bg": "#E67E22", "fg": "#FFFFFF", "hex": "E67E22"},
    "Medium":   {"bg": "#F1C40F", "fg": "#2C3E50", "hex": "F1C40F"},
    "Low":      {"bg": "#27AE60", "fg": "#FFFFFF", "hex": "27AE60"},
}

PRIORITY_BADGE = {
    "Immediate":  '<span style="background:#C0392B;color:#fff;padding:2px 8px;border-radius:4px;font-size:0.8em">⚠ IMMEDIATE</span>',
    "Short-term": '<span style="background:#E67E22;color:#fff;padding:2px 8px;border-radius:4px;font-size:0.8em">SHORT-TERM</span>',
    "Long-term":  '<span style="background:#27AE60;color:#fff;padding:2px 8px;border-radius:4px;font-size:0.8em">LONG-TERM</span>',
}


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _risk_from_row(row: dict) -> dict:
    risk = row.get("risk", {})
    if isinstance(risk, dict):
        return risk
    # dataclass fallback
    return vars(risk) if hasattr(risk, "__dict__") else {}


def _list_to_html(items: list) -> str:
    if not items:
        return "<em>None</em>"
    return "<ul>" + "".join(f"<li>{i}</li>" for i in items) + "</ul>"


def _list_to_str(items: list, sep: str = "; ") -> str:
    return sep.join(str(i) for i in items) if items else ""


def _ts() -> str:
    return datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")


# ══════════════════════════════════════════════════════════════════════════════
# 1. JSON Export
# ══════════════════════════════════════════════════════════════════════════════

def to_json(rows: list[dict], summary: dict, meta: dict,
            output_path: Optional[Path] = None) -> str:
    """
    Serialise full analysis to JSON string and optionally write to file.
    """
    payload = {
        "generated_at": _ts(),
        "metadata": meta,
        "summary": summary,
        "hazop_table": rows,
    }
    json_str = json.dumps(payload, indent=2, ensure_ascii=False, default=str)

    if output_path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json_str, encoding="utf-8")
        logger.info(f"JSON saved → {output_path}")

    return json_str


# ══════════════════════════════════════════════════════════════════════════════
# 2. HTML Export
# ══════════════════════════════════════════════════════════════════════════════

_HTML_STYLE = """
<style>
  body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 13px;
         color: #2C3E50; background: #F8F9FA; margin: 20px; }
  h1   { color: #1A252F; border-bottom: 3px solid #2980B9; padding-bottom: 8px; }
  h2   { color: #2471A3; margin-top: 30px; }
  .meta-box { background: #EAF2F8; border-left: 5px solid #2980B9;
              padding: 12px 16px; margin-bottom: 24px; border-radius: 4px; }
  .summary-grid { display: flex; gap: 16px; flex-wrap: wrap; margin-bottom: 24px; }
  .summary-card { background: #fff; border: 1px solid #D5D8DC; border-radius: 8px;
                  padding: 14px 20px; min-width: 150px; text-align: center;
                  box-shadow: 0 1px 4px rgba(0,0,0,0.08); }
  .summary-card .num { font-size: 2em; font-weight: 700; }
  .card-critical { border-top: 4px solid #C0392B; }
  .card-high     { border-top: 4px solid #E67E22; }
  .card-medium   { border-top: 4px solid #F1C40F; }
  .card-low      { border-top: 4px solid #27AE60; }
  table  { width: 100%; border-collapse: collapse; background: #fff;
           box-shadow: 0 1px 4px rgba(0,0,0,0.08); margin-bottom: 32px; }
  thead  { background: #1A252F; color: #fff; }
  th     { padding: 10px 12px; text-align: left; font-weight: 600;
           border: 1px solid #2C3E50; white-space: nowrap; }
  td     { padding: 8px 10px; border: 1px solid #D5D8DC;
           vertical-align: top; font-size: 12px; }
  tr:nth-child(even) td { background: #FDFEFE; }
  tr:hover td { background: #EBF5FB; }
  .risk-critical { background: #C0392B !important; color: #fff;
                   font-weight: 700; text-align: center; border-radius: 4px; }
  .risk-high     { background: #E67E22 !important; color: #fff;
                   font-weight: 700; text-align: center; border-radius: 4px; }
  .risk-medium   { background: #F1C40F !important; color: #2C3E50;
                   font-weight: 700; text-align: center; border-radius: 4px; }
  .risk-low      { background: #27AE60 !important; color: #fff;
                   font-weight: 700; text-align: center; border-radius: 4px; }
  ul { margin: 0; padding-left: 16px; }
  li { margin-bottom: 2px; }
  .node-label { background: #2980B9; color: #fff; border-radius: 4px;
                padding: 1px 6px; font-size: 0.85em; font-weight: 600; }
  .footer { color: #7F8C8D; font-size: 11px; margin-top: 32px;
            border-top: 1px solid #D5D8DC; padding-top: 8px; }
  @media print {
    body { margin: 0; font-size: 10px; }
    .summary-grid { display: block; }
  }
</style>
"""

_HAZOP_COLUMNS = [
    ("Node",            "node_id"),
    ("Equipment",       "equipment"),
    ("Chemical",        "chemical"),
    ("Parameter",       "parameter"),
    ("Guide Word",      "guide_word"),
    ("Deviation",       "deviation"),
    ("Causes",          "causes"),
    ("Consequences",    "consequences"),
    ("Existing Safeguards", "safeguards_existing"),
    ("Recommended Safeguards", "safeguards_recommended"),
    ("Risk Level",      "_risk_level"),
    ("Risk Score",      "_risk_score"),
    ("Actions",         "actions"),
    ("Priority",        "action_priority"),
    ("Historical Ref",  "historical_ref"),
]


def _render_cell(col_key: str, row: dict) -> str:
    """Render a table cell as HTML."""
    risk = _risk_from_row(row)

    if col_key == "_risk_level":
        level = risk.get("risk_level", "Low")
        cls   = f"risk-{level.lower()}"
        return f'<td class="{cls}">{level}</td>'

    if col_key == "_risk_score":
        score = risk.get("risk_score", 0)
        level = risk.get("risk_level", "Low")
        cls   = f"risk-{level.lower()}"
        return f'<td class="{cls}" style="text-align:center">{score}/25</td>'

    if col_key == "node_id":
        val = row.get(col_key, "")
        return f'<td><span class="node-label">{val}</span></td>'

    if col_key == "action_priority":
        pri   = row.get(col_key, "Long-term")
        badge = PRIORITY_BADGE.get(pri, pri)
        return f"<td>{badge}</td>"

    val = row.get(col_key, "")
    if isinstance(val, list):
        return f"<td>{_list_to_html(val)}</td>"

    return f"<td>{val}</td>"


def to_html(rows: list[dict], summary: dict, meta: dict,
            output_path: Optional[Path] = None) -> str:
    """
    Generate a complete, print-ready HTML HAZOP worksheet.
    """
    source = meta.get("source", "Incident Report")
    doc_id = meta.get("doc_id", "")
    date   = meta.get("date",   _ts())

    rb = summary.get("risk_breakdown", {})
    pb = summary.get("priority_breakdown", {})
    total = summary.get("total_rows", len(rows))

    # ── Summary cards ─────────────────────────────────────────────────────────
    summary_html = f"""
    <div class="summary-grid">
      <div class="summary-card card-critical">
        <div class="num" style="color:#C0392B">{rb.get('Critical', 0)}</div>
        <div>Critical</div>
      </div>
      <div class="summary-card card-high">
        <div class="num" style="color:#E67E22">{rb.get('High', 0)}</div>
        <div>High</div>
      </div>
      <div class="summary-card card-medium">
        <div class="num" style="color:#b7950b">{rb.get('Medium', 0)}</div>
        <div>Medium</div>
      </div>
      <div class="summary-card card-low">
        <div class="num" style="color:#27AE60">{rb.get('Low', 0)}</div>
        <div>Low</div>
      </div>
      <div class="summary-card">
        <div class="num">{total}</div>
        <div>Total Rows</div>
      </div>
      <div class="summary-card">
        <div class="num">{pb.get('Immediate', 0)}</div>
        <div>Immediate<br>Actions</div>
      </div>
    </div>
    """

    # ── Risk matrix legend ────────────────────────────────────────────────────
    matrix_html = _render_risk_matrix_html()

    # ── HAZOP table ───────────────────────────────────────────────────────────
    header_cells = "".join(f"<th>{col}</th>" for col, _ in _HAZOP_COLUMNS)
    table_rows_html = ""
    for row in rows:
        cells = "".join(_render_cell(key, row) for _, key in _HAZOP_COLUMNS)
        table_rows_html += f"<tr>{cells}</tr>\n"

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>HAZOP Analysis — {source}</title>
  {_HTML_STYLE}
</head>
<body>
  <h1>⚗ HAZOP Safety Analysis Report</h1>

  <div class="meta-box">
    <strong>Source:</strong> {source} &nbsp;|&nbsp;
    <strong>Doc ID:</strong> {doc_id} &nbsp;|&nbsp;
    <strong>Incident Date:</strong> {date} &nbsp;|&nbsp;
    <strong>Generated:</strong> {_ts()} &nbsp;|&nbsp;
    <strong>Nodes:</strong> {summary.get('unique_nodes', '—')} &nbsp;|&nbsp;
    <strong>Deviations:</strong> {summary.get('unique_deviations', '—')}
  </div>

  <h2>📊 Risk Summary</h2>
  {summary_html}

  <h2>🔢 Risk Matrix</h2>
  {matrix_html}

  <h2>📋 HAZOP Worksheet</h2>
  <table>
    <thead><tr>{header_cells}</tr></thead>
    <tbody>
{table_rows_html}
    </tbody>
  </table>

  <div class="footer">
    Generated by HAZOP NLP System &mdash; For engineering review only.
    Not a substitute for a formal HAZOP study conducted by qualified engineers.
    References: IEC 61882, CCPS Guidelines for Hazard Evaluation Procedures,
    U.S. Chemical Safety Board.
  </div>
</body>
</html>"""

    if output_path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(html, encoding="utf-8")
        logger.info(f"HTML saved → {output_path}")

    return html


def _render_risk_matrix_html() -> str:
    """Generate a 5×5 HTML risk matrix."""
    levels = [1, 2, 3, 4, 5]
    colour_map = {
        "Critical": "#C0392B", "High": "#E67E22",
        "Medium": "#F1C40F",   "Low":  "#27AE60",
    }
    matrix = {
        (s, l): "Critical" if s * l >= 15 else
                "High"     if s * l >= 9  else
                "Medium"   if s * l >= 4  else "Low"
        for s in levels for l in levels
    }
    rows_html = ""
    for s in reversed(levels):
        cells = f'<td style="font-weight:700;background:#1A252F;color:#fff;text-align:center">S{s}</td>'
        for l in levels:
            level = matrix[(s, l)]
            bg    = colour_map[level]
            fg    = "#fff" if level in ("Critical", "High", "Low") else "#2C3E50"
            cells += (
                f'<td style="background:{bg};color:{fg};text-align:center;'
                f'font-size:11px;padding:6px">{level}<br>{s*l}</td>'
            )
        rows_html += f"<tr>{cells}</tr>\n"

    header = '<tr><td style="background:#1A252F"></td>' + \
             "".join(f'<th style="background:#2980B9;color:#fff;text-align:center">L{l}</th>' for l in levels) + \
             "</tr>"
    return (
        f'<table style="border-collapse:collapse;width:auto;margin-bottom:16px">'
        f'<thead>{header}</thead><tbody>{rows_html}</tbody></table>'
        f'<p style="font-size:11px;color:#7F8C8D">'
        f'S = Severity (1–5) | L = Likelihood (1–5) | Score = S × L</p>'
    )


# ══════════════════════════════════════════════════════════════════════════════
# 3. Excel Export
# ══════════════════════════════════════════════════════════════════════════════

def to_excel(rows: list[dict], summary: dict, meta: dict,
             ner_result: Optional[dict] = None,
             output_path: Optional[Path] = None) -> bytes:
    """
    Generate an Excel workbook with sheets:
      1. Summary       — stats + metadata
      2. HAZOP Table   — full worksheet with colour-coded risk
      3. NER Entities  — extracted entities (if ner_result provided)
      4. Risk Matrix   — 5×5 matrix
    """
    if not _OPENPYXL:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = Workbook()

    # ── Styles ────────────────────────────────────────────────────────────────
    header_font  = Font(bold=True, color="FFFFFF", size=11)
    header_fill  = PatternFill("solid", fgColor="1A252F")
    thin_border  = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )
    wrap_align   = Alignment(wrap_text=True, vertical="top")
    center_align = Alignment(horizontal="center", vertical="center")

    risk_fills = {
        k: PatternFill("solid", fgColor=v["hex"])
        for k, v in RISK_COLOURS.items()
    }
    risk_fonts = {
        "Critical": Font(bold=True, color="FFFFFF"),
        "High":     Font(bold=True, color="FFFFFF"),
        "Medium":   Font(bold=True, color="2C3E50"),
        "Low":      Font(bold=True, color="FFFFFF"),
    }

    def _hdr(ws, row: int, col: int, val: str):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font   = header_font
        cell.fill   = header_fill
        cell.border = thin_border
        cell.alignment = center_align

    def _cell(ws, row: int, col: int, val, align=None):
        cell = ws.cell(row=row, column=col, value=val)
        cell.border = thin_border
        cell.alignment = align or wrap_align
        return cell

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"

    ws1.merge_cells("A1:D1")
    title_cell = ws1["A1"]
    title_cell.value = "HAZOP NLP Analysis — Summary"
    title_cell.font  = Font(bold=True, size=16, color="1A252F")
    title_cell.alignment = center_align

    info = [
        ("Source",          meta.get("source", "")),
        ("Doc ID",          meta.get("doc_id", "")),
        ("Incident Date",   meta.get("date", "")),
        ("Chemicals Found", ", ".join(meta.get("chemicals", []))),
        ("Generated At",    _ts()),
        ("",                ""),
        ("Total HAZOP Rows",    summary.get("total_rows", 0)),
        ("Unique Nodes",        summary.get("unique_nodes", 0)),
        ("Unique Deviations",   summary.get("unique_deviations", 0)),
        ("",                    ""),
        ("RISK BREAKDOWN",  ""),
        ("Critical",        summary.get("risk_breakdown", {}).get("Critical", 0)),
        ("High",            summary.get("risk_breakdown", {}).get("High", 0)),
        ("Medium",          summary.get("risk_breakdown", {}).get("Medium", 0)),
        ("Low",             summary.get("risk_breakdown", {}).get("Low", 0)),
        ("",                ""),
        ("ACTION PRIORITIES", ""),
        ("Immediate",       summary.get("priority_breakdown", {}).get("Immediate", 0)),
        ("Short-term",      summary.get("priority_breakdown", {}).get("Short-term", 0)),
        ("Long-term",       summary.get("priority_breakdown", {}).get("Long-term", 0)),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws1.cell(row=i, column=1, value=k).font = Font(bold=bool(k and not v == ""))
        ws1.cell(row=i, column=2, value=v)

    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 50

    # ── Sheet 2: HAZOP Table ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("HAZOP Table")

    cols = [
        ("Node ID",          12), ("Equipment",     18), ("Chemical",      16),
        ("Parameter",        12), ("Guide Word",    12), ("Deviation",      20),
        ("Causes",           35), ("Consequences",  35),
        ("Existing Safeguards", 35), ("Recommended Safeguards", 35),
        ("Risk Level",       14), ("Risk Score",    10), ("Severity",       10),
        ("Likelihood",       10), ("Actions",       45), ("Priority",       14),
        ("Historical Ref",   35), ("Notes",         30),
    ]

    for ci, (hdr, width) in enumerate(cols, 1):
        _hdr(ws2, 1, ci, hdr)
        ws2.column_dimensions[get_column_letter(ci)].width = width

    ws2.row_dimensions[1].height = 22
    ws2.freeze_panes = "A2"

    for ri, row in enumerate(rows, 2):
        risk = _risk_from_row(row)
        level = risk.get("risk_level", "Low")

        values = [
            row.get("node_id", ""),
            row.get("equipment", ""),
            row.get("chemical", ""),
            row.get("parameter", ""),
            row.get("guide_word", ""),
            row.get("deviation", ""),
            _list_to_str(row.get("causes", [])),
            _list_to_str(row.get("consequences", [])),
            _list_to_str(row.get("safeguards_existing", [])),
            _list_to_str(row.get("safeguards_recommended", [])),
            level,
            risk.get("risk_score", 0),
            risk.get("severity", 0),
            risk.get("likelihood", 0),
            _list_to_str(row.get("actions", [])),
            row.get("action_priority", ""),
            row.get("historical_ref", ""),
            row.get("notes", ""),
        ]

        for ci, val in enumerate(values, 1):
            c = _cell(ws2, ri, ci, val)
            # Colour the risk level and score columns
            if ci in (11, 12):
                c.fill      = risk_fills.get(level, PatternFill())
                c.font      = risk_fonts.get(level, Font())
                c.alignment = center_align
        ws2.row_dimensions[ri].height = 55

    # ── Sheet 3: NER Entities ─────────────────────────────────────────────────
    if ner_result:
        ws3 = wb.create_sheet("NER Entities")
        ner_cols = [("Entity Text", 25), ("Label", 14), ("Label Description", 28),
                    ("Confidence", 12), ("Source", 10), ("Sentence", 60)]
        for ci, (hdr, width) in enumerate(ner_cols, 1):
            _hdr(ws3, 1, ci, hdr)
            ws3.column_dimensions[get_column_letter(ci)].width = width
        ws3.freeze_panes = "A2"
        for ri, ent in enumerate(ner_result.get("entities", []), 2):
            vals = [
                ent.get("text", ""),
                ent.get("label", ""),
                ent.get("label_desc", ""),
                ent.get("score", 0),
                ent.get("source", ""),
                ent.get("sentence", "")[:120],
            ]
            for ci, val in enumerate(vals, 1):
                _cell(ws3, ri, ci, val)

    # ── Sheet 4: Risk Matrix ──────────────────────────────────────────────────
    ws4 = wb.create_sheet("Risk Matrix")
    ws4.merge_cells("A1:G1")
    ws4["A1"].value = "5×5 Risk Matrix (Severity × Likelihood)"
    ws4["A1"].font  = Font(bold=True, size=13)

    sev_labels = {1: "Negligible", 2: "Minor", 3: "Moderate", 4: "Major", 5: "Catastrophic"}
    lik_labels = {1: "Rare", 2: "Unlikely", 3: "Possible", 4: "Likely", 5: "Almost Certain"}

    levels_range = [1, 2, 3, 4, 5]
    matrix_data  = {
        (s, l): ("Critical" if s * l >= 15 else "High" if s * l >= 9 else "Medium" if s * l >= 4 else "Low")
        for s in levels_range for l in levels_range
    }

    # Header row
    ws4.cell(row=3, column=1, value="S \\ L").font = Font(bold=True)
    for l in levels_range:
        c = ws4.cell(row=3, column=l + 1, value=f"L{l} — {lik_labels[l]}")
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2980B9")
        c.alignment = center_align
        ws4.column_dimensions[get_column_letter(l + 1)].width = 18

    for s in reversed(levels_range):
        row_i = 3 + (6 - s)
        label_cell = ws4.cell(row=row_i, column=1, value=f"S{s} — {sev_labels[s]}")
        label_cell.font = Font(bold=True, color="FFFFFF")
        label_cell.fill = PatternFill("solid", fgColor="1A252F")
        ws4.column_dimensions["A"].width = 22
        for l in levels_range:
            level = matrix_data[(s, l)]
            c = ws4.cell(row=row_i, column=l + 1, value=f"{level}\n({s*l})")
            c.fill      = risk_fills[level]
            c.font      = risk_fonts[level]
            c.alignment = center_align
            c.border    = thin_border
        ws4.row_dimensions[row_i].height = 30

    # ── Save ──────────────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    excel_bytes = buf.read()

    if output_path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(excel_bytes)
        logger.info(f"Excel saved → {output_path}")

    return excel_bytes


# ══════════════════════════════════════════════════════════════════════════════
# 4. CSV Export
# ══════════════════════════════════════════════════════════════════════════════

CSV_FIELDS = [
    "node_id", "equipment", "chemical", "parameter", "guide_word", "deviation",
    "causes", "consequences", "safeguards_existing", "safeguards_recommended",
    "risk_level", "risk_score", "severity", "likelihood",
    "actions", "action_priority", "historical_ref", "notes",
]


def to_csv(rows: list[dict], output_path: Optional[Path] = None) -> str:
    """
    Export HAZOP rows to CSV string, optionally writing to file.
    """
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_FIELDS, extrasaction="ignore",
                            lineterminator="\n")
    writer.writeheader()

    for row in rows:
        risk  = _risk_from_row(row)
        flat  = {
            **{k: row.get(k, "") for k in CSV_FIELDS},
            "causes":                  _list_to_str(row.get("causes", [])),
            "consequences":            _list_to_str(row.get("consequences", [])),
            "safeguards_existing":     _list_to_str(row.get("safeguards_existing", [])),
            "safeguards_recommended":  _list_to_str(row.get("safeguards_recommended", [])),
            "actions":                 _list_to_str(row.get("actions", [])),
            "risk_level":              risk.get("risk_level", ""),
            "risk_score":              risk.get("risk_score", ""),
            "severity":                risk.get("severity", ""),
            "likelihood":              risk.get("likelihood", ""),
        }
        writer.writerow(flat)

    csv_str = buf.getvalue()

    if output_path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(csv_str, encoding="utf-8")
        logger.info(f"CSV saved → {output_path}")

    return csv_str


# ══════════════════════════════════════════════════════════════════════════════
# Unified Export — write all formats at once
# ══════════════════════════════════════════════════════════════════════════════

def export_all(
    rows: list[dict],
    summary: dict,
    meta: dict,
    ner_result: Optional[dict] = None,
    doc_id: str = "analysis",
    output_dir: Optional[Path] = None,
) -> dict[str, Path]:
    """
    Write JSON, HTML, Excel, and CSV to output_dir.
    Returns dict of format → path.
    """
    output_dir = Path(output_dir) if output_dir else OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)

    paths: dict[str, Path] = {}

    paths["json"] = output_dir / f"{doc_id}_hazop.json"
    to_json(rows, summary, meta, paths["json"])

    paths["html"] = output_dir / f"{doc_id}_hazop.html"
    to_html(rows, summary, meta, paths["html"])

    paths["csv"] = output_dir / f"{doc_id}_hazop.csv"
    to_csv(rows, paths["csv"])

    if _OPENPYXL:
        paths["excel"] = output_dir / f"{doc_id}_hazop.xlsx"
        to_excel(rows, summary, meta, ner_result, paths["excel"])
    else:
        logger.warning("openpyxl not installed — skipping Excel export.")

    logger.info(f"All exports written to {output_dir}")
    return paths


# ══════════════════════════════════════════════════════════════════════════════
# CLI
# ══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    # Quick smoke test using engine demo data
    import sys
    sys.path.insert(0, str(Path(__file__).parent.parent))
    from src.hazop_engine import HAZOPEngine, summarise_hazop, rows_to_dicts

    sample = (
        "A hydrogen leak from a corroded heat exchanger at the ethylene plant caused "
        "an explosion. The pressure relief valve failed to open due to corrosion. "
        "The reactor temperature rose rapidly following loss of cooling water. "
        "Two operators suffered burns. Recommended: install gas detectors and upgrade PRVs."
    )
    engine = HAZOPEngine()
    rows   = engine.run_from_text(sample, use_bert=False)
    dicts  = rows_to_dicts(rows)
    summ   = summarise_hazop(rows)
    meta   = {"source": "smoke_test", "doc_id": "test001", "date": "2024-01-01", "chemicals": ["hydrogen"]}

    paths  = export_all(dicts, summ, meta, doc_id="smoke_test")
    for fmt, p in paths.items():
        print(f"  {fmt:6s} → {p}")
